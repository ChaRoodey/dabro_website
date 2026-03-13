from io import BytesIO
from typing import Any
from openpyxl import load_workbook

from app.core.config import settings


def _norm_header(head: Any) -> str:
    return str(head).strip() if head is not None else ''


def _to_int(value: Any, field: str, row: int, errors: list[dict]) -> int | None:
    if value is None or str(value).strip() == '':
        errors.append({
            "row": row,
            "field": field,
            "message": f"Поле '{field}' не может быть пустым",
        })
        return None

    try:
        if isinstance(value, float):
            if not value.is_integer():
                errors.append({
                    "row": row,
                    "field": field,
                    "message": f"Поле '{field}' должно быть целым числом",
                })
                return None
            return int(value)

        return int(str(value).strip())
    except Exception:
        errors.append({
            "row": row,
            "field": field,
            "message": f"Поле '{field}' должно быть числом, вместо этого {value!r}",
        })
        return None


def _to_str_or_none(value):
    if value is None:
        return None
    value = str(value).strip()
    return value if value else None


def parse_products_sheet(file_bytes: bytes) -> dict:
    wb = load_workbook(filename=BytesIO(file_bytes), data_only=True)

    if settings.SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Лист '{settings.SHEET_NAME}' не найден. Доступные листы: {wb.sheetnames}")

    ws = wb[settings.SHEET_NAME]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    headers = [_norm_header(header) for header in header_row]
    col_index = {name: idx for idx, name in enumerate(headers) if name}

    missing = settings.REQUIRED_COLUMNS - set(col_index.keys())
    if missing:
        raise ValueError(f"Нет обязательных колонок: {sorted(missing)}")

    items: list[dict] = []
    errors: list[dict] = []

    for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell is None or str(cell).strip() == '' for cell in row):
            continue

        row_errors: list[dict] = []

        def get(col: str):
            i = col_index[col]
            return row[i] if i < len(row) else None

        excel_product_id = _to_int(get('id'), 'id', r_idx, row_errors)
        brand = _to_str_or_none(get('Брэнд'))
        category = _to_str_or_none(get('Категория'))
        description = _to_str_or_none(get('Описание'))
        size = _to_str_or_none(get('Объем'))
        title = _to_str_or_none(get('Название'))
        cost = _to_int(get('Стоимость'), 'Стоимость', r_idx, row_errors)
        items_left = _to_int(get('Остаток'), 'Остаток', r_idx, row_errors)

        if row_errors:
            errors.extend(row_errors)
            continue

        items.append({
            "row": r_idx,
            "excel_product_id": excel_product_id,
            "brand": brand,
            "category": category,
            "description": description,
            "size": size,
            "title": title,
            "cost": cost,
            "items_left": items_left,
        })

    return {
        "items": items,
        "errors": errors,
    }
